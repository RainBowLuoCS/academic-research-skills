"""Generate the concise seven-slide editable VCM presentation."""

from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from generate_vcm_ppt import (
    BLUE,
    CYAN,
    GREEN,
    H,
    INK,
    LIGHT,
    LIGHT_BLUE,
    MID,
    NAVY,
    ORANGE,
    PALE_BLUE,
    RED,
    W,
    WHITE,
    add_box,
    add_callout,
    add_circle,
    add_footer,
    add_line,
    add_native_chart,
    add_rich_text,
    add_text,
    add_title,
    add_token_grid,
    new_slide,
)


OUT_FILE = Path(__file__).resolve().parent / "VCM_7页精简汇报_可编辑版.pptx"


def build_deck():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    prs.core_properties.title = "VCM：7页精简论文汇报"
    prs.core_properties.subject = "Editable concise research presentation"
    prs.core_properties.author = "Generated from Luo et al. (2025)"
    prs.core_properties.comments = "All diagrams and charts use editable PowerPoint objects."

    # 1 — Title and one-sentence contribution
    s = new_slide(prs, bg=NAVY)
    add_text(s, "VCM", 0.78, 1.02, 2.10, 0.72, size=39, bold=True, color=WHITE)
    add_text(
        s,
        "Vision Concept Modeling with Adaptive Vision Token\nCompression via Instruction Fine-Tuning",
        0.80,
        1.78,
        8.65,
        1.35,
        size=25,
        bold=True,
        color=WHITE,
    )
    add_text(
        s,
        "让 LVLM 根据任务指令，动态决定“看多少、看哪里”",
        0.82,
        3.46,
        8.55,
        0.52,
        size=20,
        bold=True,
        color=CYAN,
    )
    add_token_grid(s, 9.80, 1.25, 8, 6, active={11, 12, 19, 20, 27}, cell=0.25, gap=0.06)
    add_line(s, 10.85, 3.35, 10.85, 4.05, color=WHITE, width=2.3, arrow=True)
    for i, (label, color) in enumerate([("概念 1", BLUE), ("概念 2", GREEN), ("概念 K", ORANGE)]):
        add_circle(s, 9.62 + i * 1.10, 4.20, 0.80, fill=color, line=WHITE)
        add_text(s, label, 9.66 + i * 1.10, 4.42, 0.72, 0.25, size=9, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_box(s, 0.82, 5.28, 8.45, 0.72, fill=NAVY, line=CYAN, radius=False, line_width=1.4)
    add_text(
        s,
        "核心贡献：把视觉压缩从固定 token reduction，升级为 instruction-conditioned concept modeling",
        1.04,
        5.42,
        8.02,
        0.42,
        size=16,
        bold=True,
        color=WHITE,
        align=PP_ALIGN.CENTER,
    )
    add_text(s, "Luo et al. · arXiv:2504.19627v2 · 2025", 0.82, 6.62, 5.5, 0.28, size=10, color=WHITE)
    add_text(s, "01", 12.03, 6.55, 0.55, 0.30, size=10, color=WHITE, align=PP_ALIGN.RIGHT)

    # 2 — Problem and definition
    s = new_slide(prs)
    add_title(s, "问题：LVLM 处理所有视觉 token，但任务通常只需要少量概念")
    add_text(s, "当前范式", 0.82, 1.30, 2.0, 0.36, size=17, bold=True, color=MID, align=PP_ALIGN.CENTER)
    add_token_grid(s, 1.05, 1.84, 10, 7, cell=0.22, gap=0.05)
    add_text(s, "576+ tokens", 1.40, 4.06, 2.25, 0.35, size=16, bold=True, color=RED, align=PP_ALIGN.CENTER)
    add_line(s, 4.05, 2.86, 5.03, 2.86, color=RED, width=2.4, arrow=True)
    add_box(s, 5.12, 1.84, 2.20, 2.10, fill=PALE_BLUE, line=BLUE)
    add_text(s, "LLM", 5.58, 2.19, 1.30, 0.45, size=27, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    add_text(s, "FLOPs ↑\nLatency ↑\nMemory ↑", 5.60, 2.78, 1.25, 0.85, size=14, bold=True, color=RED, align=PP_ALIGN.CENTER)
    add_line(s, 7.46, 2.86, 8.40, 2.86, color=GREEN, width=2.4, arrow=True)
    add_text(s, "VCM", 7.65, 2.36, 0.65, 0.30, size=12, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
    add_box(s, 8.52, 1.46, 3.75, 2.92, fill=WHITE, line=GREEN, line_width=1.7)
    add_text(s, "Vision Concept Model", 8.86, 1.75, 3.05, 0.42, size=21, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
    add_text(s, "由 instruction 动态决定", 9.18, 2.38, 2.42, 0.32, size=14, color=MID, align=PP_ALIGN.CENTER)
    add_rich_text(
        s,
        [("数量 ", INK, False), ("K", RED, True), ("  +  空间位置", INK, False)],
        9.08,
        2.84,
        2.62,
        0.45,
        size=18,
        align=PP_ALIGN.CENTER,
    )
    for i, col in enumerate([BLUE, GREEN, ORANGE]):
        add_circle(s, 9.18 + i * 0.90, 3.52, 0.55, fill=col, line=WHITE)
    add_callout(s, "关键区别：VCM 不只是删 token，而是保留具有语义和空间对应关系的 concepts")
    add_footer(s, 2)

    # 3 — Evidence and technical challenges
    s = new_slide(prs)
    add_title(s, "出发点：文本先验可以估计每个样本需要多少视觉信息")
    panels = [
        ("Response keywords ↑", "所需 vision length ↑", BLUE, False),
        ("Instruction keywords ↑", "所需 vision length ↓", ORANGE, True),
        ("Δ keywords ↑", "更稳定的负相关", NAVY, True),
    ]
    for i, (head, foot, color, down) in enumerate(panels):
        x = 0.72 + i * 4.18
        add_box(s, x, 1.30, 3.72, 2.90, fill=PALE_BLUE if i == 2 else WHITE, line=color)
        add_text(s, head, x + 0.25, 1.56, 3.20, 0.35, size=16, bold=True, color=color, align=PP_ALIGN.CENTER)
        add_line(s, x + 0.65, 3.48, x + 3.10, 3.48, color=MID)
        add_line(s, x + 0.65, 3.48, x + 0.65, 2.12, color=MID)
        if down:
            add_line(s, x + 0.98, 2.30, x + 2.82, 3.22, color=color, width=2.8)
        else:
            add_line(s, x + 0.98, 3.20, x + 2.82, 2.30, color=color, width=2.8)
        add_text(s, foot, x + 0.52, 3.66, 2.70, 0.30, size=11, bold=i == 2, color=color, align=PP_ALIGN.CENTER)
    challenges = [
        ("挑战 1", "无 concept-level 标注", "Semantic alignment\n+ keyword selection", BLUE),
        ("挑战 2", "concept length 动态变化", "Forward–backward\noptimization", ORANGE),
    ]
    for i, (tag, problem, solution, color) in enumerate(challenges):
        x = 1.18 + i * 6.20
        add_box(s, x, 4.68, 5.55, 1.05, fill=WHITE, line=color)
        add_text(s, tag, x + 0.18, 4.86, 1.00, 0.32, size=13, bold=True, color=color, align=PP_ALIGN.CENTER)
        add_text(s, problem, x + 1.25, 4.78, 1.80, 0.48, size=14, bold=True, color=INK, align=PP_ALIGN.CENTER)
        add_line(s, x + 3.08, 5.20, x + 3.48, 5.20, color=color, arrow=True)
        add_text(s, solution, x + 3.55, 4.76, 1.75, 0.52, size=12, bold=True, color=color, align=PP_ALIGN.CENTER)
    add_callout(s, "方法逻辑：文本先验提供“目标长度”，动态规划学习“具体位置”")
    add_footer(s, 3, source="Paper Fig. 2 and Sections 3.1–3.3")

    # 4 — Complete framework
    s = new_slide(prs)
    add_title(s, "方法：四步完成 instruction-conditioned vision concept modeling")
    stages = [
        ("1", "Semantic\nAlignment", "让 selector 学会\n图文语义相关性", BLUE),
        ("2", "Keyword\nSelection", "K > E(K)\n选 image-related 词", ORANGE),
        ("3", "Dynamic\nLength", "文本先验映射为\n目标长度 L", GREEN),
        ("4", "Forward–\nBackward", "学习 token 位置\n并合并连续片段", NAVY),
    ]
    for i, (num, head, body, color) in enumerate(stages):
        x = 0.62 + i * 3.18
        add_box(s, x, 1.40, 2.68, 3.95, fill=PALE_BLUE if i in (0, 2) else WHITE, line=color, line_width=1.5)
        add_circle(s, x + 0.92, 1.72, 0.80, fill=color, line=WHITE)
        add_text(s, num, x + 1.10, 1.91, 0.44, 0.25, size=17, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_text(s, head, x + 0.28, 2.68, 2.12, 0.82, size=19, bold=True, color=color, align=PP_ALIGN.CENTER)
        add_text(s, body, x + 0.28, 3.74, 2.12, 0.85, size=14, color=INK, align=PP_ALIGN.CENTER)
        if i < 3:
            add_line(s, x + 2.70, 3.35, x + 3.12, 3.35, color=MID, width=1.8, arrow=True)
    add_box(s, 2.18, 5.66, 8.98, 0.50, fill=LIGHT_BLUE, line=NAVY, radius=False)
    add_text(
        s,
        "Image + Instruction → adaptive Hᵥᶜ → LLM response",
        2.45,
        5.72,
        8.44,
        0.34,
        size=18,
        bold=True,
        color=NAVY,
        align=PP_ALIGN.CENTER,
    )
    add_callout(s, "训练不需要昂贵的 concept annotation；推理只把少量任务相关 concepts 送入 LLM")
    add_footer(s, 4, source="Editable reconstruction of the paper's Figure 3")

    # 5 — Core algorithm
    s = new_slide(prs)
    add_title(s, "核心算法：目标长度估计 + Forward–Backward 路径优化")
    add_box(s, 0.72, 1.22, 5.80, 2.12, fill=PALE_BLUE, line=BLUE)
    add_text(s, "目标长度", 1.03, 1.46, 1.40, 0.38, size=18, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
    add_text(s, "L = ⌊ M · S · (1 − Norm(Nₖₑᵧ)) ⌋", 1.12, 2.08, 5.05, 0.52, size=23, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    add_text(s, "告诉模型“保留多少”", 2.05, 2.72, 3.10, 0.30, size=13, color=MID, align=PP_ALIGN.CENTER)
    add_box(s, 6.82, 1.22, 5.80, 2.12, fill=WHITE, line=GREEN)
    add_text(s, "扩展目标序列", 7.13, 1.46, 1.80, 0.38, size=18, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
    add_text(s, "∅  ★  ∅  ★  ∅", 8.15, 2.05, 3.20, 0.48, size=26, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
    add_text(s, "告诉模型“哪些位置可组成 concepts”", 8.02, 2.72, 3.48, 0.30, size=13, color=MID, align=PP_ALIGN.CENTER)
    add_text(s, "α(t,l)", 1.10, 4.18, 1.35, 0.46, size=24, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
    add_line(s, 2.55, 4.42, 4.25, 4.42, color=BLUE, width=2.2, arrow=True)
    add_box(s, 4.38, 3.77, 4.55, 1.34, fill=LIGHT_BLUE, line=NAVY)
    add_text(s, "p(Zᵥ|Yᵥ) = Σₗ α(t,l) · β(t,l)", 4.70, 4.00, 3.92, 0.38, size=19, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    add_text(s, "Lᵥ꜀ₘ = −log p(Zᵥ|Yᵥ)", 5.18, 4.48, 2.95, 0.32, size=15, color=RED, align=PP_ALIGN.CENTER)
    add_line(s, 9.05, 4.42, 10.73, 4.42, color=GREEN, width=2.2, arrow=True)
    add_text(s, "β(t,l)", 10.83, 4.18, 1.35, 0.46, size=24, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
    add_box(s, 3.78, 5.43, 5.78, 0.53, fill=WHITE, line=ORANGE)
    add_text(s, "搜索复杂度：O(2ᴹ) → O(M²)，随后 Segment Merging", 4.02, 5.51, 5.30, 0.32, size=15, bold=True, color=ORANGE, align=PP_ALIGN.CENTER)
    add_callout(s, "训练最大化所有合法路径；推理选取最可能路径并合并相邻保留 tokens")
    add_footer(s, 5, source="Paper Sections 3.3–3.4 and Appendix E")

    # 6 — Results in two editable charts
    s = new_slide(prs)
    add_title(s, "结果：更少的计算量，同时保持或提升任务性能")
    chart1 = add_native_chart(
        s,
        ["LLaVA\n576", "VCM\n128", "VCM\n64"],
        [("FLOPs (T)", [4.62, 1.71, 1.24])],
        0.62,
        1.34,
        5.55,
        3.90,
        colors=[BLUE],
        ymax=5.0,
        legend=False,
    )
    chart1.value_axis.has_title = True
    chart1.value_axis.axis_title.text_frame.text = "FLOPs (T)"
    chart1.value_axis.axis_title.text_frame.paragraphs[0].font.size = Pt(9)
    chart2 = add_native_chart(
        s,
        ["FastV\n192", "PDrop\n192", "Sparse\n192", "VisionZip\n192", "LLaVA\n576", "VCM\n144"],
        [("VQA Avg. (%)", [53.1, 57.8, 57.9, 59.1, 59.5, 60.8])],
        6.52,
        1.34,
        6.18,
        3.90,
        colors=[ORANGE],
        ymin=50,
        ymax=62,
        legend=False,
    )
    chart2.series[0].points[5].format.fill.solid()
    chart2.series[0].points[5].format.fill.fore_color.rgb = RED
    chart2.series[0].points[5].format.line.color.rgb = RED
    add_text(s, "计算效率", 2.38, 1.12, 2.00, 0.30, size=14, bold=True, color=MID, align=PP_ALIGN.CENTER)
    add_text(s, "11-benchmark VQA 平均表现", 8.20, 1.12, 2.80, 0.30, size=14, bold=True, color=MID, align=PP_ALIGN.CENTER)
    add_box(s, 0.88, 5.42, 3.32, 0.54, fill=PALE_BLUE, line=BLUE)
    add_text(s, "最高约 85% FLOPs reduction", 1.04, 5.50, 3.00, 0.32, size=14, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
    add_box(s, 4.98, 5.42, 3.32, 0.54, fill=WHITE, line=GREEN)
    add_text(s, "dense perception 同步增强", 5.14, 5.50, 3.00, 0.32, size=14, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
    add_box(s, 9.10, 5.42, 3.32, 0.54, fill=WHITE, line=RED)
    add_text(s, "144 tokens 获得最高 VQA Avg.", 9.26, 5.50, 3.00, 0.32, size=14, bold=True, color=RED, align=PP_ALIGN.CENTER)
    add_callout(s, "VCM 的收益不仅是压缩：它还学习到更具表示力的视觉概念")
    add_footer(s, 6, source="Paper Table 1, Table 10 and Appendix D")

    # 7 — Conclusion and limitations
    s = new_slide(prs)
    add_title(s, "结论：VCM 将视觉压缩重新定义为任务条件化概念建模")
    takeaways = [
        ("问题", "全部 vision tokens 带来高计算冗余", RED),
        ("方法", "文本先验 + 动态长度 + Forward–Backward", BLUE),
        ("结果", "最高约 85% FLOPs reduction，性能保持或提升", GREEN),
    ]
    for i, (head, body, color) in enumerate(takeaways):
        y = 1.30 + i * 1.20
        add_circle(s, 0.90, y, 0.68, fill=color, line=WHITE)
        add_text(s, str(i + 1), 1.05, y + 0.15, 0.38, 0.24, size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_text(s, head, 1.88, y + 0.04, 1.15, 0.52, size=20, bold=True, color=color, align=PP_ALIGN.CENTER)
        add_text(s, body, 3.28, y + 0.04, 8.58, 0.52, size=19, bold=i == 2, color=INK)
    add_box(s, 0.92, 5.04, 11.48, 0.84, fill=PALE_BLUE, line=ORANGE)
    add_text(s, "局限", 1.17, 5.23, 1.10, 0.36, size=17, bold=True, color=ORANGE, align=PP_ALIGN.CENTER)
    add_text(
        s,
        "关键词可能有偏差；长度估计仍较粗；concept 可解释性需要更直接的定量评估",
        2.48,
        5.18,
        9.50,
        0.46,
        size=15,
        color=INK,
        align=PP_ALIGN.CENTER,
    )
    add_box(s, 1.25, 6.20, 10.82, 0.58, fill=NAVY, line=NAVY, radius=False)
    add_text(
        s,
        "Take-home：让模型不再“看完再回答”，而是“按问题选择该看的概念”",
        1.52,
        6.29,
        10.28,
        0.38,
        size=17,
        bold=True,
        color=WHITE,
        align=PP_ALIGN.CENTER,
    )
    add_footer(s, 7, source="Paper conclusion and limitations")

    prs.save(OUT_FILE)
    return OUT_FILE


if __name__ == "__main__":
    print(build_deck())
