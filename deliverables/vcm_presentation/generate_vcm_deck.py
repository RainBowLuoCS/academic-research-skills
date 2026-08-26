"""Build the VCM research deck: seven slides, editable, one design system.

Typography, spacing, colour and chart styling are defined once at the top and
reused by every slide, so the deck reads as a single designed artefact rather
than a collection of boxes. Diagrams and charts are native PowerPoint objects;
only two paper figures are placed as images, because they are photographic
evidence that cannot be redrawn as shapes.

    python3 generate_vcm_deck.py
"""

from __future__ import annotations

from pathlib import Path

import pymupdf
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
OUT_FILE = ROOT / "VCM_7页精简汇报_可编辑版.pptx"
PAPER = Path(
    "/home/ubuntu/.cursor/projects/workspace/uploads/"
    "2504.19627v2_compressed_d67d.pdf"
)

# ---------------------------------------------------------------- design system

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

INK = RGBColor(0x16, 0x21, 0x2B)
NAVY = RGBColor(0x12, 0x3A, 0x5F)
BLUE = RGBColor(0x1F, 0x6F, 0xB2)
SKY = RGBColor(0x8F, 0xBD, 0xE0)
TINT = RGBColor(0xED, 0xF4, 0xFA)
CORAL = RGBColor(0xC8, 0x44, 0x3C)
AMBER = RGBColor(0xD9, 0x8A, 0x2B)
TEAL = RGBColor(0x2E, 0x8B, 0x84)
GREY = RGBColor(0x6B, 0x78, 0x85)
MUTED = RGBColor(0xAA, 0xB4, 0xBE)
HAIRLINE = RGBColor(0xDD, 0xE4, 0xEA)
PAGE = RGBColor(0xFF, 0xFF, 0xFF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

LATIN = "Segoe UI"
CJK = "Microsoft YaHei"

# Horizontal grid: one left margin, one content width, used by every slide.
ML = 0.85
CW = 11.63
CR = ML + CW

# Vertical rhythm.
Y_EYEBROW = 0.44
Y_TITLE = 0.76
Y_RULE = 1.46
Y_BODY = 1.82
Y_BAND = 6.32
Y_FOOT = 7.03


def _set_font(run, size, bold, color, spacing=None, latin=LATIN, cjk=CJK):
    font = run.font
    font.size = Pt(size)
    font.bold = bold
    font.color.rgb = color
    font.name = latin
    rpr = run._r.get_or_add_rPr()
    for tag, typeface in (("a:latin", latin), ("a:ea", cjk), ("a:cs", latin)):
        existing = rpr.find(qn(tag))
        if existing is not None:
            rpr.remove(existing)
        el = rpr.makeelement(qn(tag), {"typeface": typeface})
        rpr.append(el)
    if spacing:
        rpr.set("spc", str(int(spacing * 100)))


def text(
    slide,
    body,
    x,
    y,
    w,
    h,
    *,
    size=16,
    bold=False,
    color=INK,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.MIDDLE,
    spacing=None,
    line_spacing=None,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = valign
    for i, line in enumerate(body.split("\n")):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.alignment = align
        if line_spacing:
            para.line_spacing = line_spacing
        run = para.add_run()
        run.text = line
        _set_font(run, size, bold, color, spacing)
    return box


def rich(slide, parts, x, y, w, h, *, size=16, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = frame.paragraphs[0]
    para.alignment = align
    for body, color, bold in parts:
        run = para.add_run()
        run.text = body
        _set_font(run, size, bold, color)
    return box


def _flat(shape):
    shape.shadow.inherit = False
    return shape


def panel(slide, x, y, w, h, *, fill=WHITE, line=HAIRLINE, width=0.75, radius=0.04):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    try:
        shape.adjustments[0] = radius
    except (IndexError, KeyError):
        pass
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(width)
    return _flat(shape)


def square(slide, x, y, w, h, *, fill, line=None, width=0.5):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(width)
    return _flat(shape)


def dot(slide, cx, cy, d, *, fill, line=None, width=0.75):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.OVAL,
        Inches(cx - d / 2),
        Inches(cy - d / 2),
        Inches(d),
        Inches(d),
    )
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(width)
    return _flat(shape)


def rule(slide, x1, y1, x2, y2, *, color=HAIRLINE, width=0.75, arrow=False):
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    conn.line.color.rgb = color
    conn.line.width = Pt(width)
    if arrow:
        ln = conn.line._get_or_add_ln()
        tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
        ln.append(tail)
    return conn


def heading(slide, eyebrow, title, *, accent=BLUE):
    text(
        slide,
        eyebrow,
        ML,
        Y_EYEBROW,
        CW,
        0.26,
        size=10.5,
        bold=True,
        color=accent,
        spacing=1.6,
    )
    text(slide, title, ML, Y_TITLE, CW, 0.62, size=27, bold=True, color=INK)
    rule(slide, ML, Y_RULE, ML + 1.15, Y_RULE, color=accent, width=3.0)


def band(slide, message, *, accent=BLUE):
    panel(slide, ML, Y_BAND, CW, 0.56, fill=TINT, line=None, radius=0.10)
    square(slide, ML, Y_BAND, 0.055, 0.56, fill=accent)
    text(
        slide,
        message,
        ML + 0.36,
        Y_BAND,
        CW - 0.72,
        0.56,
        size=15.5,
        bold=True,
        color=NAVY,
    )


def footer(slide, page, source):
    text(slide, source, ML, Y_FOOT, CW - 0.9, 0.22, size=8, color=MUTED)
    text(
        slide,
        f"{page:02d}",
        CR - 0.7,
        Y_FOOT - 0.02,
        0.7,
        0.26,
        size=10,
        bold=True,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )


def blank(prs, *, bg=PAGE):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = bg
    return slide


# ------------------------------------------------------------------- chart look


def bar_chart(
    slide,
    categories,
    values,
    x,
    y,
    w,
    h,
    *,
    color=BLUE,
    highlight=None,
    highlight_color=CORAL,
    number_format="0.00",
    label_size=12,
    gap_width=58,
    ymin=None,
    ymax=None,
):
    data = CategoryChartData()
    data.categories = categories
    data.add_series("series", values)
    frame = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
        data,
    )
    chart = frame.chart
    chart.has_title = False
    chart.has_legend = False
    chart.font.size = Pt(11)
    chart.font.name = LATIN
    chart.font.color.rgb = GREY

    value_axis = chart.value_axis
    value_axis.visible = False
    value_axis.has_major_gridlines = False
    if ymin is not None:
        value_axis.minimum_scale = ymin
    if ymax is not None:
        value_axis.maximum_scale = ymax

    category_axis = chart.category_axis
    category_axis.has_major_gridlines = False
    category_axis.format.line.color.rgb = HAIRLINE
    category_axis.format.line.width = Pt(0.75)
    ticks = category_axis.tick_labels
    ticks.font.size = Pt(11)
    ticks.font.name = LATIN
    ticks.font.color.rgb = GREY

    plot = chart.plots[0]
    plot.gap_width = gap_width
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format = number_format
    labels.number_format_is_linked = False
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    labels.font.size = Pt(label_size)
    labels.font.bold = True
    labels.font.name = LATIN
    labels.font.color.rgb = INK

    series = chart.series[0]
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = color
    series.format.line.fill.background()
    if highlight is not None:
        point = series.points[highlight]
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = highlight_color
        point.format.line.fill.background()
    return chart


# ------------------------------------------------------------------ paper crops


def crop(name, page_index, rect, zoom=5.0):
    ASSETS.mkdir(exist_ok=True)
    target = ASSETS / name
    if not target.exists():
        doc = pymupdf.open(PAPER)
        pix = doc[page_index].get_pixmap(
            matrix=pymupdf.Matrix(zoom, zoom),
            clip=pymupdf.Rect(*rect),
            alpha=False,
        )
        pix.save(target)
    return target


def picture(slide, path, x, y, w):
    shape = slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))
    shape.shadow.inherit = False
    return shape


# ----------------------------------------------------------------------- slides


def slide_title(prs):
    s = blank(prs, bg=NAVY)
    # Restrained motif: three tints of navy, sparse, right-aligned.
    tints = [RGBColor(0x1B, 0x4E, 0x77), RGBColor(0x24, 0x62, 0x8F), BLUE]
    for row in range(4):
        for col in range(6):
            keep = (row, col) in {(1, 2), (1, 3), (2, 2), (2, 3), (2, 4), (3, 4)}
            square(
                s,
                9.62 + col * 0.52,
                2.28 + row * 0.52,
                0.40,
                0.40,
                fill=tints[2] if keep else tints[(row + col) % 2],
            )
    text(
        s,
        "PAPER PRESENTATION · arXiv:2504.19627v2",
        ML,
        1.62,
        7.6,
        0.28,
        size=10.5,
        bold=True,
        color=SKY,
        spacing=1.8,
    )
    text(s, "VCM", ML, 2.06, 7.6, 1.02, size=68, bold=True, color=WHITE)
    text(
        s,
        "Vision Concept Modeling with Adaptive\nVision Token Compression",
        ML,
        3.20,
        7.9,
        1.06,
        size=24,
        bold=True,
        color=WHITE,
        line_spacing=1.18,
    )
    rule(s, ML, 4.52, ML + 1.15, 4.52, color=SKY, width=3.0)
    text(
        s,
        "让模型按指令决定看多少、看哪里",
        ML,
        4.76,
        7.9,
        0.44,
        size=19,
        color=SKY,
    )
    text(
        s,
        "Luo, Shan, Chen, Liu, Wang, Yang, Xia · SIAT · NUS · 2025",
        ML,
        6.42,
        7.9,
        0.28,
        size=11,
        color=RGBColor(0x7E, 0xA6, 0xC6),
    )
    return s


def slide_problem(prs):
    s = blank(prs)
    heading(s, "问题 · PROBLEM", "整幅图都进入 LLM，代价随序列长度增长")

    panel(s, ML, Y_BODY, 5.28, 3.30, fill=WHITE, line=HAIRLINE)
    text(s, "w/o VCM", ML + 0.42, 2.10, 2.2, 0.30, size=12, bold=True, color=GREY)
    text(s, "576", ML + 0.42, 2.46, 3.0, 1.10, size=76, bold=True, color=INK)
    text(
        s,
        "vision tokens",
        ML + 0.46,
        3.62,
        3.0,
        0.30,
        size=13,
        color=GREY,
    )
    for row in range(4):
        for col in range(6):
            square(
                s,
                ML + 3.42 + col * 0.28,
                2.42 + row * 0.28,
                0.22,
                0.22,
                fill=SKY if (row + col) % 2 else BLUE,
            )
    text(
        s,
        "每个 patch 都参与注意力计算",
        ML + 0.42,
        4.34,
        4.5,
        0.30,
        size=14,
        color=INK,
    )

    rule(s, 6.38, 3.47, 7.02, 3.47, color=CORAL, width=2.5, arrow=True)
    text(
        s,
        "instruction\naware",
        6.26,
        2.82,
        0.90,
        0.56,
        size=10,
        bold=True,
        color=CORAL,
        align=PP_ALIGN.CENTER,
        line_spacing=1.12,
    )

    panel(s, 7.20, Y_BODY, 5.28, 3.30, fill=TINT, line=None)
    text(s, "w/ VCM", 7.62, 2.10, 2.2, 0.30, size=12, bold=True, color=BLUE)
    text(s, "16", 7.62, 2.46, 3.0, 1.10, size=76, bold=True, color=NAVY)
    text(s, "vision tokens", 7.66, 3.62, 3.0, 0.30, size=13, color=GREY)
    for row in range(4):
        for col in range(6):
            keep = (row, col) in {(1, 2), (1, 3), (2, 2), (2, 3)}
            square(
                s,
                10.62 + col * 0.28,
                2.42 + row * 0.28,
                0.22,
                0.22,
                fill=BLUE if keep else RGBColor(0xD8, 0xE4, 0xEF),
            )
    rich(
        s,
        [
            ("同样答对 ", INK, False),
            ("“What color is the dog's collar?”", CORAL, True),
        ],
        7.62,
        4.34,
        4.6,
        0.30,
        size=14,
    )

    for i, (value, label) in enumerate(
        [("4.62 T", "FLOPs, 576 tokens"), ("57.8 ms", "latency, 576 tokens"), ("1/8", "序列长度缩放")]
    ):
        x = ML + i * 3.98
        rule(s, x, 5.42, x + 3.50, 5.42, color=HAIRLINE)
        text(s, value, x, 5.56, 2.0, 0.44, size=22, bold=True, color=INK)
        text(s, label, x + 2.02, 5.62, 1.60, 0.32, size=11.5, color=GREY)

    band(s, "瓶颈不在模型能力，而在“无差别处理全部视觉 token”")
    footer(s, 2, "Paper Figure 1, Table 10 and Appendix D")
    return s


def slide_insight(prs, strip):
    s = blank(prs)
    heading(s, "洞察 · INSIGHT", "需要多少视觉信息，文本先验已经透露")

    findings = [
        ("↑", "Response keywords", "回答涉及的图像关键词越多，所需视觉长度越长", BLUE),
        ("↓", "Instruction keywords", "问题越具体，所需视觉长度越短", AMBER),
        ("↓", "Response − Instruction", "两者之差噪声更低，负相关更稳定", TEAL),
    ]
    for i, (glyph, term, note, color) in enumerate(findings):
        y = Y_BODY + i * 1.05
        text(s, glyph, ML, y, 0.42, 0.62, size=30, bold=True, color=color, align=PP_ALIGN.CENTER)
        text(s, term, ML + 0.52, y + 0.02, 3.85, 0.30, size=15, bold=True, color=INK)
        text(s, note, ML + 0.52, y + 0.34, 4.05, 0.30, size=12.5, color=GREY)
        if i < 2:
            rule(s, ML, y + 0.86, ML + 4.45, y + 0.86, color=HAIRLINE)

    panel(s, ML, 5.06, 4.45, 0.86, fill=TINT, line=None)
    text(
        s,
        "L = ⌊ M · S · (1 − Norm(Nₖₑᵧ)) ⌋",
        ML + 0.24,
        5.06,
        3.97,
        0.50,
        size=17,
        bold=True,
        color=NAVY,
    )
    text(
        s,
        "粗粒度指令数据即可自监督估计概念预算",
        ML + 0.24,
        5.52,
        3.97,
        0.28,
        size=11.5,
        color=GREY,
    )

    labels = ["no mask", "person, yellow", "person", "all masked"]
    picture(s, strip, 5.72, 2.22, 6.76)
    for i, label in enumerate(labels):
        text(
            s,
            label,
            5.78 + i * 1.69,
            1.86,
            1.60,
            0.28,
            size=11.5,
            bold=True,
            color=CORAL if i in (1, 2) else GREY,
            align=PP_ALIGN.CENTER,
        )
    rule(s, 5.72, 4.62, 12.48, 4.62, color=HAIRLINE)
    text(
        s,
        "遮蔽的关键词越多，保留的视觉概念越稀疏 —— 概念选择确实由指令驱动",
        5.72,
        4.76,
        6.76,
        0.32,
        size=13,
        color=INK,
    )
    text(
        s,
        "Q: where is the person in yellow?     A: the center of the image.",
        5.72,
        5.20,
        6.76,
        0.30,
        size=12,
        color=GREY,
    )

    band(s, "文本先验决定“保留多少”，动态规划再决定“保留哪里”")
    footer(s, 3, "Paper Figure 2 and Figure 4 (top row)")
    return s


def slide_method(prs):
    s = blank(prs)
    heading(s, "方法 · METHOD", "三步构建 instruction-conditioned vision concepts")

    columns = [
        ("A", "Semantic alignment", "对齐图像、文本与 LLM 语义空间", BLUE),
        ("B", "Keyword selection", "自适应筛出与图像相关的关键词", AMBER),
        ("C", "Concept modeling", "学习可变长度概念及其位置", TEAL),
    ]
    for i, (tag, title, note, color) in enumerate(columns):
        x = ML + i * 3.98
        panel(s, x, Y_BODY, 3.50, 3.62, fill=WHITE if i != 1 else TINT, line=HAIRLINE)
        square(s, x, Y_BODY, 3.50, 0.055, fill=color)
        text(s, tag, x + 0.30, 2.02, 0.5, 0.42, size=20, bold=True, color=color)
        text(s, title, x + 0.30, 2.50, 2.90, 0.34, size=17, bold=True, color=INK)
        text(s, note, x + 0.30, 2.86, 2.90, 0.30, size=12.5, color=GREY)

        if i == 0:
            for j, (label, tone) in enumerate(
                [("Gⱽ", BLUE), ("Gᵀ", AMBER), ("Gᴸᴸᴹ", NAVY)]
            ):
                dot(s, x + 0.85 + j * 0.92, 3.86, 0.72, fill=None, line=tone, width=1.4)
                text(
                    s,
                    label,
                    x + 0.49 + j * 0.92,
                    3.70,
                    0.72,
                    0.32,
                    size=12,
                    bold=True,
                    color=tone,
                    align=PP_ALIGN.CENTER,
                )
            rule(s, x + 1.21, 3.86, x + 1.41, 3.86, color=HAIRLINE)
            rule(s, x + 2.13, 3.86, x + 2.33, 3.86, color=HAIRLINE)
            text(
                s,
                "Lₛₐ  语义对齐损失",
                x + 0.30,
                4.62,
                2.90,
                0.32,
                size=13,
                bold=True,
                color=BLUE,
            )
        elif i == 1:
            words = [("person", True), ("yellow", True), ("where", False), ("the", False)]
            for j, (word, hot) in enumerate(words):
                bx = x + 0.30 + (j % 2) * 1.52
                by = 3.48 + (j // 2) * 0.46
                panel(
                    s,
                    bx,
                    by,
                    1.38,
                    0.36,
                    fill=CORAL if hot else WHITE,
                    line=None if hot else HAIRLINE,
                    radius=0.14,
                )
                text(
                    s,
                    word,
                    bx,
                    by,
                    1.38,
                    0.36,
                    size=12,
                    bold=hot,
                    color=WHITE if hot else MUTED,
                    align=PP_ALIGN.CENTER,
                )
            text(
                s,
                "K > E(K)  阈值以上视为关键词",
                x + 0.30,
                4.62,
                2.90,
                0.32,
                size=13,
                bold=True,
                color=AMBER,
            )
        else:
            for row in range(4):
                for col in range(7):
                    keep = (row, col) in {(1, 2), (1, 3), (2, 2), (2, 3), (2, 4)}
                    square(
                        s,
                        x + 0.34 + col * 0.42,
                        3.42 + row * 0.28,
                        0.34,
                        0.22,
                        fill=TEAL if keep else RGBColor(0xE4, 0xE9, 0xED),
                    )
            text(
                s,
                "Lᵥ꜀ₘ  动态概念损失",
                x + 0.30,
                4.62,
                2.90,
                0.32,
                size=13,
                bold=True,
                color=TEAL,
            )
        if i < 2:
            rule(s, x + 3.60, 3.63, x + 3.88, 3.63, color=MUTED, width=1.4, arrow=True)

    rule(s, ML, 5.66, CR, 5.66, color=HAIRLINE)
    rich(
        s,
        [
            ("总目标  ", GREY, False),
            ("L = L", INK, True),
            ("ₙₜₚ", INK, True),
            (" + ε(r) · L", INK, True),
            ("ᵥ꜀ₘ", INK, True),
            ("        无需 concept-level 标注", GREY, False),
        ],
        ML,
        5.78,
        CW,
        0.36,
        size=16,
    )

    band(s, "两阶段训练：先学会“什么与图像相关”，再学会“保留多少、保留哪里”")
    footer(s, 4, "Editable reconstruction of Paper Figure 3")
    return s


def slide_algorithm(prs):
    s = blank(prs)
    heading(s, "算法 · ALGORITHM", "Forward–Backward 学习可变长度对齐")

    panel(s, ML, Y_BODY, 4.30, 3.62, fill=TINT, line=None)
    text(s, "目标序列", ML + 0.32, 2.00, 3.66, 0.30, size=12, bold=True, color=BLUE)
    text(
        s,
        "Zᵥ = [ ∅, ★, ∅, ★, ∅ ]",
        ML + 0.32,
        2.34,
        3.66,
        0.46,
        size=21,
        bold=True,
        color=NAVY,
    )
    text(
        s,
        "长度 L 由文本先验给出，位置未知",
        ML + 0.32,
        2.82,
        3.66,
        0.28,
        size=12,
        color=GREY,
    )
    rule(s, ML + 0.32, 3.24, ML + 3.98, 3.24, color=HAIRLINE)
    text(s, "优化目标", ML + 0.32, 3.42, 3.66, 0.30, size=12, bold=True, color=BLUE)
    text(
        s,
        "p(Zᵥ|Yᵥ) = Σₗ α(t,l) · β(t,l)",
        ML + 0.32,
        3.74,
        3.66,
        0.40,
        size=17,
        bold=True,
        color=INK,
    )
    text(
        s,
        "Lᵥ꜀ₘ = − log p(Zᵥ|Yᵥ)",
        ML + 0.32,
        4.20,
        3.66,
        0.36,
        size=17,
        bold=True,
        color=CORAL,
    )
    rich(
        s,
        [("复杂度  ", GREY, False), ("O(2ᴹ)", MUTED, True), ("  →  ", GREY, False), ("O(M²)", NAVY, True)],
        ML + 0.32,
        4.78,
        3.66,
        0.36,
        size=16,
    )

    # Lattice: possible alignments in light blue, optimal path in coral.
    text(s, "extended target  l", 6.05, 1.92, 2.4, 0.26, size=11, color=GREY)
    ox, oy = 6.05, 2.42
    dx, dy = 0.86, 0.55
    cols, rows = 7, 5
    path = [0, 1, 1, 2, 2, 3, 4]
    for c in range(cols - 1):
        for r in range(rows):
            rule(s, ox + c * dx, oy + r * dy, ox + (c + 1) * dx, oy + r * dy, color=RGBColor(0xE3, 0xEC, 0xF4), width=0.75)
            if r < rows - 1:
                rule(
                    s,
                    ox + c * dx,
                    oy + r * dy,
                    ox + (c + 1) * dx,
                    oy + (r + 1) * dy,
                    color=RGBColor(0xE3, 0xEC, 0xF4),
                    width=0.75,
                )
    for c in range(cols - 1):
        rule(
            s,
            ox + c * dx,
            oy + path[c] * dy,
            ox + (c + 1) * dx,
            oy + path[c + 1] * dy,
            color=CORAL,
            width=2.75,
        )
    for c in range(cols):
        for r in range(rows):
            active = r == path[c]
            dot(
                s,
                ox + c * dx,
                oy + r * dy,
                0.17 if active else 0.13,
                fill=CORAL if active else WHITE,
                line=CORAL if active else SKY,
                width=0.9,
            )
    text(
        s,
        "vision tokens  t = 1 … M",
        ox,
        oy + (rows - 1) * dy + 0.22,
        3.0,
        0.26,
        size=11,
        color=GREY,
    )
    rule(s, 11.42, 2.46, 11.78, 2.46, color=RGBColor(0xC9, 0xD9, 0xE7), width=1.6)
    text(s, "可行路径", 11.86, 2.32, 1.10, 0.28, size=11, color=GREY)
    rule(s, 11.42, 2.84, 11.78, 2.84, color=CORAL, width=2.75)
    text(s, "最优路径", 11.86, 2.70, 1.10, 0.28, size=11, bold=True, color=CORAL)

    # Segment merging strip.
    probs = [0.12, 0.84, 0.77, 0.16, 0.88, 0.73, 0.11]
    text(s, "Segment merging", ox, 5.34, 2.4, 0.28, size=12, bold=True, color=TEAL)
    for i, p in enumerate(probs):
        square(
            s,
            ox + i * 0.50,
            5.70,
            0.40,
            0.40,
            fill=BLUE if p > 0.5 else RGBColor(0xE4, 0xE9, 0xED),
        )
    rule(s, ox + 3.62, 5.90, ox + 4.08, 5.90, color=TEAL, width=1.6, arrow=True)
    for i, tone in enumerate([BLUE, TEAL]):
        dot(s, ox + 4.44 + i * 0.74, 5.90, 0.56, fill=tone)
        text(
            s,
            f"C{i + 1}",
            ox + 4.16 + i * 0.74,
            5.76,
            0.56,
            0.28,
            size=11,
            bold=True,
            color=WHITE,
            align=PP_ALIGN.CENTER,
        )
    text(
        s,
        "相邻保留 token 按概率加权\n合并为一个视觉概念",
        ML,
        5.58,
        4.30,
        0.62,
        size=12.5,
        color=GREY,
        line_spacing=1.2,
    )

    footer(s, 5, "Paper Sections 3.3–3.4, Figure 6 and Appendix E")
    return s


def slide_results(prs):
    s = blank(prs)
    heading(s, "结果 · RESULTS", "计算量大幅下降，平均性能不降反升")

    text(s, "FLOPs (T)", ML, Y_BODY, 2.6, 0.28, size=12, bold=True, color=GREY)
    bar_chart(
        s,
        ["LLaVA-1.5\n576", "VCM\n128", "VCM\n64"],
        [4.62, 1.71, 1.24],
        ML - 0.28,
        2.14,
        4.50,
        3.02,
        color=SKY,
        highlight=1,
        highlight_color=BLUE,
        number_format="0.00",
        ymax=5.4,
    )
    text(
        s,
        "理论最高约 85% FLOPs reduction",
        ML,
        5.22,
        4.10,
        0.30,
        size=12.5,
        color=INK,
    )

    rule(s, 5.62, 2.20, 5.62, 5.12, color=HAIRLINE)

    text(s, "11-benchmark VQA average (%)", 5.94, Y_BODY, 4.2, 0.28, size=12, bold=True, color=GREY)
    bar_chart(
        s,
        ["FastV", "PDrop", "SparseVLM", "VisionZip", "LLaVA-1.5", "VCM"],
        [53.1, 57.8, 57.9, 59.1, 59.5, 60.8],
        5.70,
        2.14,
        6.78,
        3.02,
        color=RGBColor(0xC7, 0xD5, 0xE2),
        highlight=5,
        highlight_color=CORAL,
        number_format="0.0",
        ymin=50,
        ymax=63,
        gap_width=48,
    )
    rich(
        s,
        [
            ("VCM 仅用 ", INK, False),
            ("144", CORAL, True),
            (" tokens，其余方法为 192 或 576", INK, False),
        ],
        5.94,
        5.22,
        6.54,
        0.30,
        size=12.5,
    )

    for i, (value, label, tone) in enumerate(
        [("+1.3", "Avg. vs. LLaVA-1.5", CORAL), ("+5.7", "SEED", BLUE), ("+4.9", "VizWiz", TEAL)]
    ):
        x = ML + i * 3.98
        rule(s, x, 5.68, x + 3.50, 5.68, color=HAIRLINE)
        text(s, value, x, 5.80, 1.35, 0.42, size=22, bold=True, color=tone)
        text(s, label, x + 1.32, 5.86, 2.15, 0.32, size=11.5, color=GREY)

    band(s, "更少的 token 得到更高的平均分：学到的是表示力，而不只是压缩率")
    footer(s, 6, "Paper Table 1, Table 10 and Appendix D")
    return s


def slide_conclusion(prs, triplet):
    s = blank(prs)
    heading(s, "结论 · CONCLUSION", "从 token compression 到 vision concept modeling")

    rows = [
        ("贡献", "形式化 vision concept model:\n由指令决定概念数量与空间位置", BLUE),
        ("方法", "语义对齐 + 文本先验长度估计\n+ forward–backward 对齐", NAVY),
        ("证据", "约 85% FLOPs reduction\nVQA 平均分反超 576-token baseline", TEAL),
        ("局限", "关键词偏差、长度估计偏粗\n概念可解释性缺少直接度量", CORAL),
    ]
    for i, (tag, body, tone) in enumerate(rows):
        y = Y_BODY + i * 1.02
        square(s, ML, y + 0.10, 0.055, 0.56, fill=tone)
        text(s, tag, ML + 0.26, y + 0.14, 0.80, 0.44, size=15, bold=True, color=tone)
        text(
            s,
            body,
            ML + 1.24,
            y + 0.04,
            4.80,
            0.66,
            size=13.5,
            color=INK,
            line_spacing=1.25,
        )
        if i < 3:
            rule(s, ML, y + 0.86, ML + 6.34, y + 0.86, color=HAIRLINE)

    panel(s, 7.06, Y_BODY - 0.06, 5.42, 3.98, fill=TINT, line=None)
    text(
        s,
        "K-Means on CLIP ViT dense features",
        7.38,
        1.98,
        4.78,
        0.28,
        size=11.5,
        bold=True,
        color=NAVY,
        align=PP_ALIGN.CENTER,
    )
    picture(s, triplet, 7.44, 2.42, 4.66)
    for i, label in enumerate(["Image", "w/o VCM", "w/ VCM"]):
        text(
            s,
            label,
            7.48 + i * 1.56,
            4.22,
            1.48,
            0.28,
            size=11.5,
            bold=i == 2,
            color=NAVY if i == 2 else GREY,
            align=PP_ALIGN.CENTER,
        )
    rule(s, 7.44, 4.62, 12.10, 4.62, color=RGBColor(0xD5, 0xE2, 0xEE))
    text(
        s,
        "稠密特征聚类更贴合物体边界",
        7.44,
        4.76,
        4.66,
        0.30,
        size=13,
        color=INK,
        align=PP_ALIGN.CENTER,
    )
    text(
        s,
        "效率提升与细粒度感知增强可以同时发生",
        7.44,
        5.14,
        4.66,
        0.30,
        size=13,
        bold=True,
        color=TEAL,
        align=PP_ALIGN.CENTER,
    )

    panel(s, ML, Y_BAND, CW, 0.56, fill=NAVY, line=None, radius=0.10)
    rich(
        s,
        [
            ("Take-home:  ", SKY, True),
            ("不要看完再回答，而是按问题选择该看的概念", WHITE, True),
        ],
        ML + 0.36,
        Y_BAND,
        CW - 0.72,
        0.56,
        size=16,
        align=PP_ALIGN.CENTER,
    )
    footer(s, 7, "Paper Figure 8, conclusion and limitations")
    return s


def build_deck():
    strip = crop("strip_instruction_sparsity.png", 5, (165, 94, 532, 156))
    triplet = crop("strip_kmeans_triplet.png", 5, (365, 160, 531, 221))

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    prs.core_properties.title = "VCM 论文汇报"
    prs.core_properties.author = "Generated from Luo et al. (2025)"
    prs.core_properties.comments = (
        "Native editable shapes, connectors, text and charts; two paper figures "
        "are placed as photographic evidence."
    )

    slide_title(prs)
    slide_problem(prs)
    slide_insight(prs, strip)
    slide_method(prs)
    slide_algorithm(prs)
    slide_results(prs)
    slide_conclusion(prs, triplet)

    prs.save(OUT_FILE)
    return OUT_FILE


if __name__ == "__main__":
    print(build_deck())
