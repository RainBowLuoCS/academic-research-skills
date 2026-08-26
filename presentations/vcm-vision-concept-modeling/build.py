#!/usr/bin/env python3
"""Build the VCM talk: editable SVG slides, a preview PDF, and a PPTX.

The slide chrome (title, sub-line, takeaway box, citation) is described once as
a list of layout specs and then rendered twice: into SVG, and into native
PowerPoint shapes. The central figure of every slide is a standalone SVG that is
embedded in the PPTX with a PNG fallback, so PowerPoint shows it as a vector
graphic that "Convert to Shape" turns into fully native, editable shapes.

    python3 build.py            # everything
    python3 build.py --only svg # skip the PDF/PPTX steps
"""

from __future__ import annotations

import argparse
import io
import shutil
import sys
from pathlib import Path
from xml.sax.saxutils import escape

import figures
from content import CONDENSED_SLIDES, PAPER, SLIDES
from svgkit import (FILL_GREY, GREY, INK, MATH, MUTED, NAVY, RED, SANS, WHITE,
                    Canvas, _fmt, text_width)

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "out"
SLIDE_DIR = OUT / "slides"
FIG_DIR = OUT / "figures"
PNG_DIR = OUT / "png"

# --------------------------------------------------------------------- geometry

W, H = 1280, 720
PX_PER_PT = 4 / 3          # 96 dpi authoring, 72 pt-per-inch output
EMU_PER_PX = 9525

FIG_X, FIG_Y, FIG_W = 64, 146, 1152
FIG_H_SUBLINE, FIG_H_PLAIN = 356, 410

TITLE_BOX = (64, 62, 1152, 62)
SUBLINE_BOX = (64, 520, 1152, 32)
TAKE_BOX = (57, 566, 1166, 78)
CITE_BOX = (64, 682, 1152, 26)

TITLE_SIZE = 40
SUBLINE_SIZE = 17
TAKE_SIZE_MAX, TAKE_SIZE_MIN = 28, 19
CITE_SIZE = 13


# ------------------------------------------------------------ text measurement

def measure(text, size, bold=True):
    return text_width(text, size, "bold" if bold else None)


# ------------------------------------------------------------------ rich markup

def parse_rich(s):
    """``"a *b* c"`` -> ``[("a ", False), ("b", True), (" c", False)]``."""
    runs, red = [], False
    for chunk in s.split("*"):
        if chunk:
            runs.append((chunk, red))
        red = not red
    return runs


def rich_width(runs, size, bold=True):
    return sum(measure(t, size, bold) for t, _ in runs)


def wrap_rich(runs, max_w, size, bold=True):
    """Greedy word wrap that keeps the red/plain split intact."""
    words = []
    for text, red in runs:
        parts = text.split(" ")
        for i, part in enumerate(parts):
            if part:
                words.append((part, red))
            if i < len(parts) - 1:
                words.append((" ", red))
    lines, cur, cur_w = [], [], 0.0
    for word, red in words:
        w = measure(word, size, bold)
        if cur and word != " " and cur_w + w > max_w:
            while cur and cur[-1][0] == " ":
                cur.pop()
            lines.append(cur)
            cur, cur_w = [(word, red)], w
        else:
            cur.append((word, red))
            cur_w += w
    if cur:
        while cur and cur[-1][0] == " ":
            cur.pop()
        lines.append(cur)
    return [_merge(line) for line in lines]


def _merge(runs):
    out = []
    for text, red in runs:
        if out and out[-1][1] == red:
            out[-1] = (out[-1][0] + text, red)
        else:
            out.append((text, red))
    return out


def fit_rich(s, max_w, size_max, size_min, max_lines=2, bold=True):
    """Largest size at which *s* fits in *max_lines* lines of *max_w*."""
    runs = parse_rich(s)
    size = size_max
    while size > size_min:
        if rich_width(runs, size, bold) <= max_w:
            return size, [runs]
        size -= 0.5
    size = size_max
    while size > size_min:
        lines = wrap_rich(runs, max_w, size, bold)
        if len(lines) <= max_lines and all(
                rich_width(l, size, bold) <= max_w for l in lines):
            return size, lines
        size -= 0.5
    return size_min, wrap_rich(runs, max_w, size_min, bold)


# ----------------------------------------------------------------- layout specs

def text_spec(box, lines, size, *, weight="bold", color=INK, align="l",
              family=SANS, lh=None):
    return {"kind": "text", "box": box, "lines": lines, "size": size,
            "weight": weight, "color": color, "align": align, "family": family,
            "lh": lh if lh is not None else size * 1.26}


def rect_spec(box, fill=None, stroke=None, sw=2.0, rx=0):
    return {"kind": "rect", "box": box, "fill": fill, "stroke": stroke,
            "sw": sw, "rx": rx}


def chrome_specs(slide):
    """The non-figure content of one slide, as renderer-agnostic specs."""
    if slide.get("kind") == "title":
        return [
            rect_spec((0, 0, W, H), fill=NAVY),
            text_spec((0, 196, W, 80), [[(PAPER["title_main"], False)]], 62,
                      color=WHITE, align="c"),
            text_spec((0, 286, W, 48), [[(PAPER["title_sub"], False)]], 31,
                      color="#BBD3EA", align="c"),
            rect_spec((W / 2 - 210, 374, 420, 2), fill="#5A8CBE"),
            text_spec((0, 400, W, 34), [[(PAPER["authors"], False)]], 20,
                      weight=None, color=WHITE, align="c"),
            text_spec((0, 440, W, 30), [[(PAPER["affil"], False)]], 16,
                      weight=None, color="#A9C6E2", align="c"),
            rect_spec((W / 2 - 132, 522, 264, 42), stroke="#7FA8D2", sw=1.8, rx=21),
            text_spec((W / 2 - 132, 522, 264, 42), [[(PAPER["venue"], False)]],
                      17, color=WHITE, align="c"),
        ]

    specs = [text_spec(TITLE_BOX, [[(slide["title"], False)]], TITLE_SIZE,
                       color=INK, align="l")]

    if slide.get("subline"):
        size = SUBLINE_SIZE
        while size > 12 and measure(slide["subline"], size, True) > SUBLINE_BOX[2]:
            size -= 0.5
        family = MATH if slide["key"] == "dp" else SANS
        specs.append(text_spec(SUBLINE_BOX, [[(slide["subline"], False)]], size,
                               color=INK, align="l", family=family))

    if slide.get("takeaway"):
        size, lines = fit_rich(slide["takeaway"], TAKE_BOX[2] - 56,
                              TAKE_SIZE_MAX, TAKE_SIZE_MIN)
        specs.append(rect_spec(TAKE_BOX, fill=WHITE, stroke=NAVY, sw=2.2))
        specs.append(text_spec(TAKE_BOX, lines, size, color=NAVY, align="c",
                               lh=size * 1.2))

    if slide.get("cite"):
        size = CITE_SIZE
        while size > 9 and measure(slide["cite"], size, True) > CITE_BOX[2]:
            size -= 0.5
        specs.append(text_spec(CITE_BOX, [[(slide["cite"], False)]], size,
                               color="#3C4A57", align="c"))
    return specs


def figure_box(slide):
    if slide.get("kind") == "title":
        return None
    h = FIG_H_SUBLINE if slide.get("subline") else FIG_H_PLAIN
    return (FIG_X, FIG_Y, FIG_W, h)


# --------------------------------------------------------------- SVG rendering

def draw_specs(cv: Canvas, specs):
    for sp in specs:
        x, y, w, h = sp["box"]
        if sp["kind"] == "rect":
            cv.rect(x, y, w, h, fill=sp["fill"], stroke=sp["stroke"],
                    sw=sp["sw"], rx=sp["rx"] or None)
            continue
        size, lh = sp["size"], sp["lh"]
        n = len(sp["lines"])
        base = y + h / 2 - (n - 1) * lh / 2 + size * 0.35
        for i, runs in enumerate(sp["lines"]):
            ly = base + i * lh
            # One <text> with inline tspans so the renderer does the inline
            # flow (kerning, word spacing) itself; measurement is used only to
            # place the line's left edge, because text-anchor and tspans are
            # not handled consistently across renderers.
            bold = sp["weight"] == "bold"
            line_w = sum(measure(t, size, bold) for t, _ in runs)
            lx = x + (w - line_w) / 2 if sp["align"] == "c" else x
            body = "".join(
                f'<tspan fill="{RED if red else sp["color"]}">{escape(t)}</tspan>'
                for t, red in runs)
            cv.raw(f'<text x="{_fmt(lx)}" y="{_fmt(ly)}" '
                   f'font-family="{sp["family"]}" font-size="{_fmt(size)}" '
                   + (f'font-weight="{sp["weight"]}" ' if sp["weight"] else "")
                   + f'fill="{sp["color"]}" text-anchor="start" '
                   f'xml:space="preserve">{body}</text>')


def build_figure(slide):
    """Render one slide's figure into its own Canvas, or None for the title."""
    box = figure_box(slide)
    if box is None:
        return None
    fn = getattr(figures, slide["fig"])
    cv = Canvas(box[2], box[3])
    fn(cv)
    return cv


def build_slide_svg(slide):
    cv = Canvas(W, H)
    if slide.get("kind") != "title":
        cv.rect(0, 0, W, H, fill=WHITE, stroke=None)
    draw_specs(cv, chrome_specs(slide))

    if slide.get("kind") == "title":
        motif = Canvas(W, H)
        figures.fig_title_motif(motif)
        cv.markers |= motif.markers
        cv.raw(motif.body())
    else:
        fig = build_figure(slide)
        box = figure_box(slide)
        cv.markers |= fig.markers
        cv.raw(f'<g transform="translate({_fmt(box[0])},{_fmt(box[1])})">'
               + fig.body() + "</g>")
    return cv.standalone()


# ----------------------------------------------------------------- rasterising

def to_png(svg: str, scale=2.0):
    import cairosvg
    return cairosvg.svg2png(bytestring=svg.encode("utf-8"),
                            output_width=int(W * scale),
                            output_height=int(H * scale))


def write_pdf(svgs, path):
    import cairosvg
    import pymupdf

    doc = pymupdf.open()
    for svg in svgs:
        pdf_bytes = cairosvg.svg2pdf(bytestring=svg.encode("utf-8"),
                                     output_width=W, output_height=H)
        page = pymupdf.open("pdf", pdf_bytes)
        doc.insert_pdf(page)
        page.close()
    doc.set_metadata({"title": f'{PAPER["title_main"]} — talk',
                      "author": PAPER["authors"]})
    doc.save(str(path), deflate=True)
    doc.close()


# ------------------------------------------------------------- PPTX generation

SVG_EXT_URI = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_ASVG = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"


def px(v):
    from pptx.util import Emu
    return Emu(int(round(v * EMU_PER_PX)))


def pt(v):
    from pptx.util import Pt
    return Pt(v * 0.75)


def _rgb(hex_color):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(hex_color.lstrip("#").upper())


def add_svg_picture(prs, slide, svg_bytes, png_bytes, box, name):
    """Insert a picture whose primary representation is the SVG itself."""
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.opc.package import Part
    from pptx.opc.packuri import PackURI
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import qn

    x, y, w, h = box
    pic = slide.shapes.add_picture(io.BytesIO(png_bytes), px(x), px(y),
                                   px(w), px(h))
    pic.name = name

    partname = PackURI(f"/ppt/media/{name}.svg")
    svg_part = Part(partname, "image/svg+xml", prs.part.package, svg_bytes)
    r_id = slide.part.relate_to(svg_part, RT.IMAGE)

    blip = pic._element.blipFill.find(qn("a:blip"))
    blip.append(parse_xml(
        f'<a:extLst xmlns:a="{NS_A}">'
        f'<a:ext uri="{SVG_EXT_URI}">'
        f'<asvg:svgBlip xmlns:asvg="{NS_ASVG}" xmlns:r="{NS_R}" r:embed="{r_id}"/>'
        f'</a:ext></a:extLst>'))
    return pic


def add_text_spec(slide, sp):
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN

    x, y, w, h = sp["box"]
    box = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = box.text_frame
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

    for i, runs in enumerate(sp["lines"]):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = PP_ALIGN.CENTER if sp["align"] == "c" else PP_ALIGN.LEFT
        para.line_spacing = sp["lh"] / sp["size"]
        for text, red in runs:
            run = para.add_run()
            run.text = text
            font = run.font
            font.name = "Calibri" if sp["family"] == SANS else sp["family"]
            font.size = pt(sp["size"])
            font.bold = sp["weight"] == "bold"
            font.color.rgb = _rgb(RED if red else sp["color"])
    return box


def add_rect_spec(slide, sp):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Pt as _Pt

    x, y, w, h = sp["box"]
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if sp["rx"]
                                   else MSO_SHAPE.RECTANGLE,
                                   px(x), px(y), px(w), px(h))
    shape.shadow.inherit = False
    if sp["rx"]:
        shape.adjustments[0] = min(0.5, sp["rx"] / min(w, h))
    if sp["fill"]:
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(sp["fill"])
    else:
        shape.fill.background()
    if sp["stroke"]:
        shape.line.color.rgb = _rgb(sp["stroke"])
        shape.line.width = _Pt(sp["sw"] * 0.75)
    else:
        shape.line.fill.background()
    shape.text_frame.text = ""
    return shape


def write_pptx(path, figs):
    from pptx import Presentation

    prs = Presentation()
    prs.slide_width = px(W)
    prs.slide_height = px(H)
    blank = prs.slide_layouts[6]

    for idx, slide_def in enumerate(SLIDES, start=1):
        slide = prs.slides.add_slide(blank)
        specs = chrome_specs(slide_def)

        # Background rectangles first, then the figure, then the text on top.
        for sp in specs:
            if sp["kind"] == "rect" and sp["box"][2] >= W:
                add_rect_spec(slide, sp)

        svg_bytes, png_bytes, box = figs[idx]
        add_svg_picture(prs, slide, svg_bytes, png_bytes, box,
                        f"figure{idx:02d}_{slide_def['key']}")

        for sp in specs:
            if sp["kind"] == "rect" and sp["box"][2] < W:
                add_rect_spec(slide, sp)
        for sp in specs:
            if sp["kind"] == "text":
                add_text_spec(slide, sp)

    prs.save(str(path))


# ------------------------------------------------------------------------- main

def main():
    global SLIDES, OUT, SLIDE_DIR, FIG_DIR, PNG_DIR
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--only", choices=["svg", "pdf", "pptx", "all"], default="all")
    ap.add_argument("--deck", choices=["full", "condensed"], default="full",
                    help="build the original 22-slide talk or the richer 8-slide version")
    args = ap.parse_args()

    condensed = args.deck == "condensed"
    if condensed:
        SLIDES = CONDENSED_SLIDES
        OUT = ROOT / "out-condensed"
        SLIDE_DIR = OUT / "slides"
        FIG_DIR = OUT / "figures"
        PNG_DIR = OUT / "png"
    stem = "VCM_talk_condensed" if condensed else "VCM_talk"

    for d in (SLIDE_DIR, FIG_DIR, PNG_DIR):
        if d.exists():
            shutil.rmtree(d)
        d.mkdir(parents=True, exist_ok=True)

    slide_svgs, fig_assets = [], {}
    for idx, slide_def in enumerate(SLIDES, start=1):
        name = f"{idx:02d}_{slide_def['key']}"
        svg = build_slide_svg(slide_def)
        slide_svgs.append(svg)
        (SLIDE_DIR / f"slide_{name}.svg").write_text(svg, encoding="utf-8")

        if slide_def.get("kind") == "title":
            cv = Canvas(W, H)
            figures.fig_title_motif(cv)
            box = (0, 0, W, H)
        else:
            cv = build_figure(slide_def)
            box = figure_box(slide_def)
        fig_svg = cv.standalone()
        (FIG_DIR / f"fig_{name}.svg").write_text(fig_svg, encoding="utf-8")
        fig_assets[idx] = (fig_svg.encode("utf-8"), box)

    print(f"  {len(slide_svgs)} slide SVGs -> {SLIDE_DIR.relative_to(ROOT)}")
    print(f"  {len(fig_assets)} figure SVGs -> {FIG_DIR.relative_to(ROOT)}")
    if args.only == "svg":
        return 0

    pngs = {}
    for idx, svg in enumerate(slide_svgs, start=1):
        name = f"{idx:02d}_{SLIDES[idx - 1]['key']}"
        png = to_png(svg)
        (PNG_DIR / f"slide_{name}.png").write_bytes(png)
        pngs[idx] = png

    if args.only in ("pdf", "all"):
        write_pdf(slide_svgs, OUT / f"{stem}.pdf")
        print(f"  PDF  -> {OUT.name}/{stem}.pdf")

    if args.only in ("pptx", "all"):
        import cairosvg
        figs = {}
        for idx in fig_assets:
            svg_bytes, box = fig_assets[idx]
            png = cairosvg.svg2png(bytestring=svg_bytes,
                                   output_width=int(box[2] * 2),
                                   output_height=int(box[3] * 2))
            figs[idx] = (svg_bytes, png, box)
        write_pptx(OUT / f"{stem}.pptx", figs)
        print(f"  PPTX -> {OUT.name}/{stem}.pptx")
    return 0


if __name__ == "__main__":
    sys.exit(main())
